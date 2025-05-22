"""文件处理模块"""
import os
import pandas as pd
from docx import Document
from PyPDF2 import PdfReader
import xlrd
import openpyxl
from typing import List, Dict, Any
from ..utils.text_utils import clean_for_excel

def read_text_file(file_path: str) -> str:
    """读取文本文件
    
    Args:
        file_path: 文件路径
        
    Returns:
        str: 文件内容
    """
    with open(file_path, 'r', encoding='utf-8') as f:
        return f.read()

def read_docx_file(file_path: str) -> str:
    """读取Word文档
    
    Args:
        file_path: 文件路径
        
    Returns:
        str: 文档内容
    """
    doc = Document(file_path)
    return '\n'.join([paragraph.text for paragraph in doc.paragraphs])

def read_pdf_file(file_path: str) -> str:
    """读取PDF文件
    
    Args:
        file_path: 文件路径
        
    Returns:
        str: PDF内容
    """
    reader = PdfReader(file_path)
    text = ''
    for page in reader.pages:
        text += page.extract_text() + '\n'
    return text

def read_excel_file(file_path: str) -> str:
    """读取Excel文件
    
    Args:
        file_path: 文件路径
        
    Returns:
        str: Excel内容
    """
    try:
        # 尝试使用openpyxl读取
        wb = openpyxl.load_workbook(file_path, read_only=True)
        text = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            text.append(f"Sheet: {sheet}")
            for row in ws.rows:
                text.append('\t'.join(str(cell.value) if cell.value is not None else '' for cell in row))
        return '\n'.join(text)
    except:
        # 如果失败，尝试使用xlrd读取
        wb = xlrd.open_workbook(file_path)
        text = []
        for sheet in wb.sheets():
            text.append(f"Sheet: {sheet.name}")
            for row in range(sheet.nrows):
                text.append('\t'.join(str(cell.value) if cell.value is not None else '' for cell in sheet.row(row)))
        return '\n'.join(text)

def read_file(file_path: str) -> str:
    """根据文件类型读取文件内容
    
    Args:
        file_path: 文件路径
        
    Returns:
        str: 文件内容
        
    Raises:
        ValueError: 不支持的文件类型
    """
    file_ext = os.path.splitext(file_path)[1].lower()
    
    if file_ext == '.txt':
        return read_text_file(file_path)
    elif file_ext == '.docx':
        return read_docx_file(file_path)
    elif file_ext == '.pdf':
        return read_pdf_file(file_path)
    elif file_ext in ['.xls', '.xlsx']:
        return read_excel_file(file_path)
    else:
        raise ValueError(f"不支持的文件类型: {file_ext}")

def save_qa_pairs(qa_pairs: List[Dict[str, Any]], output_file: str) -> None:
    """保存问答对到CSV文件
    
    Args:
        qa_pairs: 问答对列表
        output_file: 输出文件路径
    """
    # 创建DataFrame
    df = pd.DataFrame(qa_pairs)
    
    # 清理文本
    df['question'] = df['question'].apply(clean_for_excel)
    df['answer'] = df['answer'].apply(clean_for_excel)
    
    # 保存到CSV
    df.to_csv(output_file, index=False, encoding='utf-8-sig') 